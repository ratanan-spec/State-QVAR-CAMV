from __future__ import annotations

import argparse
import json
import platform
import sys
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy
import sklearn
import statsmodels
from openpyxl import load_workbook
from sklearn.linear_model import QuantileRegressor
from statsmodels.stats.diagnostic import het_arch
from statsmodels.tsa.stattools import adfuller


# ============================= USER SETTINGS =============================
NODES = ["EUA", "CGW", "CLEAN", "BGRN", "SPESG"]
STATE_VARIABLES = ["VIX", "US10Y"]
QUANTILES = [0.05, 0.25, 0.50, 0.75, 0.95]
ROLLING_QUANTILES = [0.05, 0.50, 0.95]
ROLLING_WINDOW = 500
ROLLING_STEP = 5
BASELINE_HORIZON = 10
ROBUSTNESS_HORIZONS = [5, 10, 20]
HIGH_STATE_PERCENTILE = 0.75
BOOTSTRAP_REPS = 2000
EXPECTED_BLOCK_LENGTH = 10
SEED = 2026
STABILITY_LIMIT = 0.995
EQUITY_CLUSTER = ["CGW", "CLEAN", "SPESG"]
POLICY_FINANCE_CLUSTER = ["EUA", "BGRN"]
# ========================================================================

ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "data" / "sustainability_markets.xlsx"
DEFAULT_OUTPUT = ROOT / "outputs"


def repair_date(value):
    """Recover the MM/DD/YYYY date convention in the original Excel export."""
    if value is None:
        return pd.NaT
    if isinstance(value, datetime):
        # Excel stored ambiguous MM/DD strings using a D/M locale.  The raw
        # workbook audit showed that swapping day and month reconstructs the
        # vendor chronology. Dates that cannot be swapped are retained.
        try:
            return pd.Timestamp(value.year, value.day, value.month)
        except ValueError:
            return pd.Timestamp(value)
    return pd.to_datetime(str(value).strip(), format="%m/%d/%Y", errors="coerce")


def load_and_align(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Sheet1" not in wb.sheetnames:
        raise ValueError("Input workbook must contain a sheet named 'Sheet1'.")
    rows = list(wb["Sheet1"].iter_rows(min_row=2, values_only=True))
    mapping = {
        "EUA": (2, 3), "CGW": (4, 5), "CLEAN": (6, 7),
        "BGRN": (8, 9), "SPESG": (10, 11), "VIX": (12, 13),
        "US10Y": (14, 15),
    }
    series, audit = {}, []
    for name, (date_col, value_col) in mapping.items():
        records = []
        raw_valid = 0
        for row in rows:
            d = repair_date(row[date_col - 1])
            v = pd.to_numeric(row[value_col - 1], errors="coerce")
            if pd.notna(d) and pd.notna(v):
                raw_valid += 1
                records.append((d, float(v)))
        frame = pd.DataFrame(records, columns=["Date", name])
        duplicates = int(frame["Date"].duplicated().sum()) if len(frame) else 0
        s = frame.drop_duplicates("Date", keep="first").set_index("Date")[name].sort_index()
        nonpositive = int((s <= 0).sum())
        if name in NODES and nonpositive:
            raise ValueError(f"{name} contains {nonpositive} nonpositive prices.")
        series[name] = s
        audit.append({
            "Series": name, "Raw_valid_rows": raw_valid,
            "Unique_dates": len(s), "Start": s.index.min(), "End": s.index.max(),
            "Duplicate_dates_removed": duplicates, "Nonpositive_values": nonpositive,
        })

    levels = pd.concat([series[x] for x in NODES + STATE_VARIABLES], axis=1, join="inner").dropna()
    if levels.empty:
        raise ValueError("No common dates remain after calendar intersection.")
    returns = 100.0 * np.log(levels[NODES]).diff().dropna()
    states = levels.loc[returns.index, STATE_VARIABLES].copy()
    return levels, returns, states, pd.DataFrame(audit)


def generalized_fevd(A: np.ndarray, sigma: np.ndarray, horizon: int):
    n = A.shape[0]
    psi = [np.linalg.matrix_power(A, h) for h in range(horizon)]
    theta = np.zeros((n, n))
    for i in range(n):
        denominator = sum((p @ sigma @ p.T)[i, i] for p in psi)
        for j in range(n):
            numerator = sum((p @ sigma)[i, j] ** 2 for p in psi) / sigma[j, j]
            theta[i, j] = numerator / denominator
    row_sums = theta.sum(axis=1, keepdims=True)
    if np.any(~np.isfinite(theta)) or np.any(row_sums <= 0):
        raise FloatingPointError("Invalid GFEVD matrix.")
    return theta / row_sums * 100.0


def cluster_components(fevd: np.ndarray):
    index = {name: i for i, name in enumerate(NODES)}
    eq = [index[x] for x in EQUITY_CLUSTER]
    pf = [index[x] for x in POLICY_FINANCE_CLUSTER]

    def directed_sum(group_a, group_b, same=False):
        total = 0.0
        for receiver in group_a:
            for source in group_b:
                if not (same and receiver == source):
                    total += fevd[receiver, source]
        return total / len(NODES)

    within_equity = directed_sum(eq, eq, same=True)
    within_policy_finance = directed_sum(pf, pf, same=True)
    cross = (directed_sum(eq, pf) + directed_sum(pf, eq))
    return within_equity, within_policy_finance, cross


def fit_qvar_connectedness(data: pd.DataFrame, tau: float, horizon: int):
    x = data.shift(1).dropna().to_numpy()
    y = data.iloc[1:].to_numpy()
    n = data.shape[1]
    A, residuals = np.zeros((n, n)), np.zeros_like(y)
    intercepts = np.zeros(n)
    for i in range(n):
        model = QuantileRegressor(
            quantile=tau, alpha=0.0, fit_intercept=True, solver="highs"
        ).fit(x, y[:, i])
        intercepts[i], A[i, :] = model.intercept_, model.coef_
        residuals[:, i] = y[:, i] - model.predict(x)

    raw_radius = float(np.max(np.abs(np.linalg.eigvals(A))))
    rescaled = raw_radius >= STABILITY_LIMIT
    if rescaled:
        A = A * (STABILITY_LIMIT / raw_radius)
    sigma = np.cov(residuals, rowvar=False, ddof=1)
    fevd = generalized_fevd(A, sigma, horizon)
    own = np.diag(fevd)
    from_ = fevd.sum(axis=1) - own
    to = fevd.sum(axis=0) - own
    net = to - from_
    tci = float(from_.sum() / n)
    within_eq, within_pf, cross = cluster_components(fevd)
    return {
        "intercepts": intercepts, "A": A, "Sigma": sigma, "FEVD": fevd,
        "FROM": from_, "TO": to, "NET": net, "TCI": tci,
        "Spectral_Radius": raw_radius, "Stability_Rescaled": rescaled,
        "Within_Equity": within_eq, "EUA_BGRN_Bilateral": within_pf,
        "Cross_Cluster": cross,
    }


def stationary_bootstrap(values, statistic=np.mean, reps=BOOTSTRAP_REPS,
                         expected_block=EXPECTED_BLOCK_LENGTH, seed=SEED):
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    if len(values) < 2:
        return np.array([np.nan] * reps)
    n, restart_probability = len(values), 1.0 / expected_block
    rng, draws = np.random.default_rng(seed), np.empty(reps)
    for b in range(reps):
        sample, idx = np.empty(n), int(rng.integers(0, n))
        for t in range(n):
            sample[t] = values[idx]
            idx = int(rng.integers(0, n)) if rng.random() < restart_probability else (idx + 1) % n
        draws[b] = statistic(sample)
    return draws


def stationary_bootstrap_group_difference(values, high, reps=BOOTSTRAP_REPS,
                                          expected_block=EXPECTED_BLOCK_LENGTH,
                                          seed=SEED):
    values, high = np.asarray(values, float), np.asarray(high, int)
    ok = np.isfinite(values)
    values, high = values[ok], high[ok]
    n, p = len(values), 1.0 / expected_block
    rng, draws = np.random.default_rng(seed), np.empty(reps)
    for b in range(reps):
        idxs, idx = np.empty(n, dtype=int), int(rng.integers(0, n))
        for t in range(n):
            idxs[t] = idx
            idx = int(rng.integers(0, n)) if rng.random() < p else (idx + 1) % n
        v, h = values[idxs], high[idxs]
        draws[b] = v[h == 1].mean() - v[h == 0].mean() if np.any(h == 1) and np.any(h == 0) else np.nan
    return draws[np.isfinite(draws)]


def diagnostics(returns: pd.DataFrame):
    rows = []
    for name in NODES:
        x = returns[name].dropna().to_numpy()
        adf = adfuller(x, maxlag=20, autolag="BIC", regression="c")
        arch = het_arch(x - x.mean(), nlags=20)
        rows.append({
            "Market": name, "ADF_statistic": adf[0], "ADF_p_value": adf[1],
            "ADF_lag_BIC": adf[2],
            "ARCH_LM_20": arch[0], "ARCH_p_value": arch[1],
        })
    return pd.DataFrame(rows)


def create_figures(outdir: Path):
    figure_dir = outdir / "figures"
    figure_dir.mkdir(exist_ok=True)
    colors = {"EUA": "#8C510A", "CGW": "#2166AC", "CLEAN": "#1B9E77",
              "BGRN": "#7570B3", "SPESG": "#D95F02"}
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9,
                         "axes.titlesize": 10, "axes.labelsize": 9})

    levels = pd.read_csv(outdir / "aligned_levels.csv", parse_dates=["Date"]).set_index("Date")
    norm = levels[NODES] / levels[NODES].iloc[0] * 100
    fig, ax = plt.subplots(figsize=(8.2, 4.5))
    for c in NODES:
        ax.plot(norm.index, norm[c], label=c, lw=1.25, color=colors[c])
    ax.set_ylabel("Normalized level (16 Mar 2020 = 100)")
    ax.legend(ncol=5, frameon=False, loc="upper left"); ax.grid(alpha=.22)
    fig.tight_layout(); fig.savefig(figure_dir / "figure1_normalized_levels.png", dpi=300); plt.close(fig)

    st = pd.read_csv(outdir / "static_tci.csv")
    dr = pd.read_csv(outdir / "directional_connectedness.csv")
    fig, axes = plt.subplots(1, 2, figsize=(8.2, 3.8), gridspec_kw={"width_ratios": [.8, 1.5]})
    axes[0].plot(st.Quantile, st.TCI, marker="o", color="#2166AC", lw=1.8)
    axes[0].set(xlabel="Quantile", ylabel="TCI (%)", title="(a) Total connectedness")
    axes[0].set_xticks(st.Quantile); axes[0].grid(alpha=.22)
    p = dr.pivot(index="Market", columns="Quantile", values="NET").loc[NODES]
    x, width = np.arange(len(NODES)), .15
    for k, q in enumerate(p.columns):
        axes[1].bar(x + (k - 2) * width, p[q], width, label=f"{q:.2f}")
    axes[1].axhline(0, color="black", lw=.7); axes[1].set_xticks(x, NODES)
    axes[1].set(ylabel="NET (percentage points)", title="(b) Net directional connectedness")
    axes[1].legend(title="Quantile", ncol=5, frameon=False, fontsize=7, title_fontsize=8)
    axes[1].grid(axis="y", alpha=.22)
    fig.tight_layout(); fig.savefig(figure_dir / "figure2_static_quantile.png", dpi=300); plt.close(fig)

    rolling = pd.read_csv(outdir / "rolling_quantile_connectedness.csv", parse_dates=["Date"])
    wide = rolling.pivot(index="Date", columns="Quantile", values="TCI")
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    for q, c, ls in [(.05, "#C00000", "-"), (.50, "#666666", "--"), (.95, "#2166AC", "-")]:
        ax.plot(wide.index, wide[q], label=f"tau = {q:.2f}", color=c, lw=1.35, ls=ls)
    ax.set_ylabel("Rolling TCI (%)"); ax.legend(ncol=3, frameon=False); ax.grid(alpha=.22)
    fig.tight_layout(); fig.savefig(figure_dir / "figure3_rolling_tci.png", dpi=300); plt.close(fig)

    vr = pd.read_csv(outdir / "vix_regime_summary.csv")
    yr = pd.read_csv(outdir / "yield_regime_summary.csv")
    fig, axes = plt.subplots(1, 2, figsize=(8.2, 3.7), sharey=True)
    for ax, df, state, title in [(axes[0], vr, "High_VIX", "VIX state"),
                                 (axes[1], yr, "High_US10Y", "U.S. 10-year yield state")]:
        piv = df.pivot(index="Quantile", columns=state, values="mean")
        xx = np.arange(len(piv))
        ax.bar(xx - .18, piv[0], .36, label="Lower state", color="#8DA0CB")
        ax.bar(xx + .18, piv[1], .36, label="High state", color="#FC8D62")
        ax.set_xticks(xx, [f"{q:.2f}" for q in piv.index]); ax.set_xlabel("Quantile")
        ax.set_title(title); ax.grid(axis="y", alpha=.22)
    axes[0].set_ylabel("Mean rolling TCI (%)"); axes[0].legend(frameon=False)
    fig.tight_layout(); fig.savefig(figure_dir / "figure4_regimes.png", dpi=300); plt.close(fig)

    # Complete networks: every off-diagonal link is drawn. Width/opacity encode weight.
    fevd = pd.read_csv(outdir / "quantile_gfevd.csv")
    pos = {"EUA": (-1, 0), "CGW": (-.3, .8), "CLEAN": (.7, .55),
           "BGRN": (.7, -.55), "SPESG": (-.3, -.8)}
    fig, axes = plt.subplots(1, 3, figsize=(8.2, 3.2))
    max_share = fevd.loc[fevd.Receiver != fevd.Shock_Source, "GFEVD_Share"].max()
    for ax, q in zip(axes, [.05, .50, .95]):
        sub = fevd[fevd.Quantile.eq(q)]
        for _, z in sub.iterrows():
            if z.Receiver == z.Shock_Source:
                continue
            x1, y1 = pos[z.Shock_Source]; x2, y2 = pos[z.Receiver]
            strength = max(float(z.GFEVD_Share) / max_share, .02)
            ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                        arrowprops=dict(arrowstyle="-|>", color="#444444",
                                        alpha=.12 + .65 * strength,
                                        lw=.25 + 2.6 * strength,
                                        shrinkA=18, shrinkB=18,
                                        connectionstyle="arc3,rad=.08"))
        for node in NODES:
            x0, y0 = pos[node]
            ax.scatter([x0], [y0], s=800, c=colors[node], edgecolors="white", zorder=3)
            ax.text(x0, y0, node, ha="center", va="center", fontsize=8, color="white", zorder=4)
        ax.set_title(f"tau = {q:.2f}"); ax.axis("off")
    fig.tight_layout(); fig.savefig(figure_dir / "figure5_complete_networks.png", dpi=300,
                                    bbox_inches="tight"); plt.close(fig)


def write_excel(outdir: Path):
    sheets = {
        "Static TCI": "static_tci.csv", "Directional": "directional_connectedness.csv",
        "GFEVD": "quantile_gfevd.csv", "Rolling": "rolling_quantile_connectedness.csv",
        "Tail contrasts": "tail_contrast_bootstrap.csv", "VIX regimes": "vix_regime_summary.csv",
        "Yield regimes": "yield_regime_summary.csv", "Joint regimes": "joint_regime_summary.csv",
        "Cluster decomposition": "cluster_regime_decomposition.csv",
        "Regime bootstrap": "regime_difference_bootstrap.csv",
        "Horizon sensitivity": "horizon_sensitivity.csv", "Diagnostics": "diagnostics.csv",
        "Descriptives": "return_descriptives.csv", "Correlations": "return_correlations.csv",
        "Data audit": "data_audit.csv", "Run notes": "run_notes.csv",
    }
    with pd.ExcelWriter(outdir / "sustainability_qvar_results.xlsx", engine="openpyxl") as writer:
        for sheet, filename in sheets.items():
            pd.read_csv(outdir / filename).to_excel(writer, sheet_name=sheet[:31], index=False)


def main(input_path: Path, outdir: Path):
    np.random.seed(SEED)
    outdir.mkdir(parents=True, exist_ok=True)
    levels, returns, states, audit = load_and_align(input_path)
    levels.to_csv(outdir / "aligned_levels.csv", index_label="Date")
    returns.to_csv(outdir / "aligned_returns.csv", index_label="Date")
    audit.to_csv(outdir / "data_audit.csv", index=False)

    desc = returns.describe(percentiles=[.05, .25, .50, .75, .95]).T
    desc.to_csv(outdir / "return_descriptives.csv", index_label="Market")
    returns.corr().to_csv(outdir / "return_correlations.csv", index_label="Market")
    diagnostics(returns).to_csv(outdir / "diagnostics.csv", index=False)

    static, directional, fevd_long, sensitivity = [], [], [], []
    for tau in QUANTILES:
        baseline = fit_qvar_connectedness(returns, tau, BASELINE_HORIZON)
        static.append({"Quantile": tau, "TCI": baseline["TCI"],
                       "Spectral_Radius": baseline["Spectral_Radius"],
                       "Stability_Rescaled": baseline["Stability_Rescaled"]})
        for i, receiver in enumerate(NODES):
            directional.append({"Quantile": tau, "Market": receiver,
                                "FROM": baseline["FROM"][i], "TO": baseline["TO"][i],
                                "NET": baseline["NET"][i]})
            for j, source in enumerate(NODES):
                fevd_long.append({"Quantile": tau, "Receiver": receiver,
                                  "Shock_Source": source,
                                  "GFEVD_Share": baseline["FEVD"][i, j]})
        for horizon in ROBUSTNESS_HORIZONS:
            result = baseline if horizon == BASELINE_HORIZON else fit_qvar_connectedness(returns, tau, horizon)
            sensitivity.append({"Quantile": tau, "Horizon": horizon, "TCI": result["TCI"]})

    pd.DataFrame(static).to_csv(outdir / "static_tci.csv", index=False)
    pd.DataFrame(directional).to_csv(outdir / "directional_connectedness.csv", index=False)
    pd.DataFrame(fevd_long).to_csv(outdir / "quantile_gfevd.csv", index=False)
    pd.DataFrame(sensitivity).to_csv(outdir / "horizon_sensitivity.csv", index=False)

    vix_cut, yield_cut = states.VIX.quantile(HIGH_STATE_PERCENTILE), states.US10Y.quantile(HIGH_STATE_PERCENTILE)
    rolling_rows = []
    ends = list(range(ROLLING_WINDOW, len(returns) + 1, ROLLING_STEP))
    for count, end in enumerate(ends, 1):
        window, date = returns.iloc[end - ROLLING_WINDOW:end], returns.index[end - 1]
        for tau in ROLLING_QUANTILES:
            result = fit_qvar_connectedness(window, tau, BASELINE_HORIZON)
            row = {"Date": date, "Quantile": tau, "TCI": result["TCI"],
                   "VIX": states.loc[date, "VIX"],
                   "High_VIX": int(states.loc[date, "VIX"] >= vix_cut),
                   "US10Y": states.loc[date, "US10Y"],
                   "High_US10Y": int(states.loc[date, "US10Y"] >= yield_cut),
                   "Spectral_Radius": result["Spectral_Radius"],
                   "Stability_Rescaled": result["Stability_Rescaled"],
                   "Within_Equity": result["Within_Equity"],
                   "EUA_BGRN_Bilateral": result["EUA_BGRN_Bilateral"],
                   "Cross_Cluster": result["Cross_Cluster"]}
            for i, name in enumerate(NODES):
                row[f"NET_{name}"] = result["NET"][i]
            rolling_rows.append(row)
        if count % 40 == 0 or count == len(ends):
            print(f"Rolling windows completed: {count}/{len(ends)}", flush=True)
    rolling = pd.DataFrame(rolling_rows)
    rolling.to_csv(outdir / "rolling_quantile_connectedness.csv", index=False)

    pivot = rolling.pivot(index="Date", columns="Quantile", values="TCI").dropna()
    contrasts = {
        "Lower_minus_Median": pivot[.05] - pivot[.50],
        "Upper_minus_Median": pivot[.95] - pivot[.50],
        "Tail_Average_minus_Median": (pivot[.05] + pivot[.95]) / 2 - pivot[.50],
    }
    pd.DataFrame(contrasts).to_csv(outdir / "rolling_tail_contrasts.csv", index_label="Date")
    contrast_rows = []
    for k, values in contrasts.items():
        draws = stationary_bootstrap(values.to_numpy(), seed=SEED)
        contrast_rows.append({"Contrast": k, "Mean": values.mean(),
                              "CI_Lower": np.quantile(draws, .025),
                              "CI_Upper": np.quantile(draws, .975),
                              "Bootstrap_Reps": BOOTSTRAP_REPS})
    pd.DataFrame(contrast_rows).to_csv(outdir / "tail_contrast_bootstrap.csv", index=False)

    def regime_summary(state):
        return rolling.groupby(["Quantile", state])["TCI"].agg(
            ["count", "mean", "median", "std", "min", "max"]).reset_index()
    regime_summary("High_VIX").to_csv(outdir / "vix_regime_summary.csv", index=False)
    regime_summary("High_US10Y").to_csv(outdir / "yield_regime_summary.csv", index=False)
    joint = rolling.groupby(["Quantile", "High_VIX", "High_US10Y"])["TCI"].agg(
        ["count", "mean", "median", "std", "min", "max"]).reset_index()
    joint["Regime"] = np.select(
        [(joint.High_VIX == 0) & (joint.High_US10Y == 0),
         (joint.High_VIX == 1) & (joint.High_US10Y == 0),
         (joint.High_VIX == 0) & (joint.High_US10Y == 1)],
        ["Low VIX / Low yield", "High VIX / Low yield", "Low VIX / High yield"],
        default="High VIX / High yield")
    joint.to_csv(outdir / "joint_regime_summary.csv", index=False)

    component_cols = ["TCI", "Within_Equity", "EUA_BGRN_Bilateral", "Cross_Cluster"]
    cluster = rolling.groupby(["Quantile", "High_VIX"])[component_cols].mean().reset_index()
    cluster.to_csv(outdir / "cluster_regime_decomposition.csv", index=False)

    regime_bootstrap = []
    for tau in ROLLING_QUANTILES:
        sub = rolling[rolling.Quantile.eq(tau)].sort_values("Date")
        for state in ["High_VIX", "High_US10Y"]:
            draws = stationary_bootstrap_group_difference(sub.TCI, sub[state], seed=SEED)
            observed = sub.loc[sub[state].eq(1), "TCI"].mean() - sub.loc[sub[state].eq(0), "TCI"].mean()
            regime_bootstrap.append({"Quantile": tau, "State": state,
                                     "High_minus_Low": observed,
                                     "CI_Lower": np.quantile(draws, .025),
                                     "CI_Upper": np.quantile(draws, .975),
                                     "Bootstrap_Reps": len(draws)})
    pd.DataFrame(regime_bootstrap).to_csv(outdir / "regime_difference_bootstrap.csv", index=False)

    notes = pd.DataFrame([
        ["Common level observations", len(levels)], ["Return observations", len(returns)],
        ["Common sample start", levels.index.min().date().isoformat()],
        ["Common sample end", levels.index.max().date().isoformat()],
        ["QVAR lag", 1], ["Baseline GFEVD horizon", BASELINE_HORIZON],
        ["Rolling window", ROLLING_WINDOW], ["Rolling step", ROLLING_STEP],
        ["VIX high-state threshold", vix_cut], ["US10Y high-state threshold", yield_cut],
        ["Bootstrap repetitions", BOOTSTRAP_REPS], ["Expected block length", EXPECTED_BLOCK_LENGTH],
        ["Random seed", SEED]], columns=["Item", "Value"])
    notes.to_csv(outdir / "run_notes.csv", index=False)

    create_figures(outdir)
    write_excel(outdir)
    manifest = {
        "input_file": input_path.name, "generated_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.version, "platform": platform.platform(),
        "packages": {"numpy": np.__version__, "pandas": pd.__version__,
                     "scipy": scipy.__version__, "scikit_learn": sklearn.__version__,
                     "statsmodels": statsmodels.__version__,
                     "matplotlib": matplotlib.__version__},
        "settings": {"nodes": NODES, "quantiles": QUANTILES,
                     "rolling_quantiles": ROLLING_QUANTILES,
                     "rolling_window": ROLLING_WINDOW, "rolling_step": ROLLING_STEP,
                     "baseline_horizon": BASELINE_HORIZON,
                     "robustness_horizons": ROBUSTNESS_HORIZONS,
                     "bootstrap_reps": BOOTSTRAP_REPS,
                     "expected_block_length": EXPECTED_BLOCK_LENGTH, "seed": SEED,
                     "adf_maxlag": 20, "adf_lag_selection": "BIC"},
        "sample": {"levels": len(levels), "returns": len(returns),
                   "start": str(levels.index.min().date()), "end": str(levels.index.max().date())},
    }
    (outdir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print("\nCompleted successfully")
    print(pd.DataFrame(static).to_string(index=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Reproduce the sustainability QVAR-GFEVD analysis.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    main(args.input.resolve(), args.output.resolve())